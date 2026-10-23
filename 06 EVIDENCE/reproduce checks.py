"""Offline reproduction of bounded checks. Original files are supplied by their owner.
Usage: python "reproduce checks.py" --sources SOURCE_DIR --out result.json [--helm HELM_EXE]
No application/container is started. No secret value is printed or saved.
Requires Python 3, openpyxl, PyYAML, beautifulsoup4. Optional Helm 3.19.0.
The source folder is matched against the delivered SHA-256 manifest, not personal names.
"""
from pathlib import Path
import argparse,hashlib,json,zipfile,tempfile,subprocess,sys
import openpyxl,yaml
from bs4 import BeautifulSoup
p=argparse.ArgumentParser();p.add_argument('--sources',required=True);p.add_argument('--out',required=True);p.add_argument('--helm');a=p.parse_args()
manifest=json.loads((Path(__file__).parent/'P01 source manifest.json').read_text(encoding='utf-8'))
expected={x['sha256']:x['id'] for x in manifest};found={}
for f in Path(a.sources).rglob('*'):
 if f.is_file():
  digest=hashlib.sha256(f.read_bytes()).hexdigest()
  if digest in expected:found[expected[digest]]=f
out={'matched_ids':sorted(found),'missing_ids':sorted(set(expected.values())-set(found)),'limits':'Local file checks only. No product runtime, scanning or rights authentication.'}
if 'E10'in found:
 soup=BeautifulSoup(found['E10'].read_bytes(),'html.parser');rows=[]
 for row in soup.find_all('tr'):
  cs=[c.get_text(' ',strip=True) for c in row.find_all(['td','th'])]
  if len(cs)>=6 and cs[0].isdigit():rows.append({'id':int(cs[0]),'severity':cs[1],'status':cs[5]})
 out['sast']=rows
if 'E03'in found:
 b=openpyxl.load_workbook(found['E03'],data_only=False);v=openpyxl.load_workbook(found['E03'],data_only=True)
 out['model']={'sheets':b.sheetnames,'formula_count':sum(c.data_type=='f' for sh in b for row in sh for c in row),'cached_errors':[(sh.title,c.coordinate,c.value) for sh in v for row in sh for c in row if c.data_type=='e'],'note':'Cached values only; live recalculation uses P05 excel recalculate.ps1.'}
if 'E13'in found:
 with tempfile.TemporaryDirectory() as td:
  td=Path(td)
  with zipfile.ZipFile(found['E13']) as z:
   for n in z.namelist():
    if not (td/n).resolve().is_relative_to(td.resolve()):raise ValueError('Unsafe ZIP member')
   z.extractall(td)
  chart=td/'pandev-metrics-distribution/helm-chart'
  comp=yaml.safe_load((td/'pandev-metrics-distribution/docker-compose/docker-compose.yml').read_text(encoding='utf-8'))
  out['compose']={k:{'image':v.get('image'),'ports':v.get('ports'),'healthcheck_present':bool(v.get('healthcheck'))} for k,v in comp['services'].items()}
  env=(td/'pandev-metrics-distribution/docker-compose/.env').read_text(encoding='utf-8');out['password_example_present']=any(x.startswith('POSTGRES_PASSWORD=') and bool(x.split('=',1)[1].strip()) for x in env.splitlines())
  if a.helm:
   tests=[]
   for args in [['version','--short'],['lint',str(chart)],['template','audit',str(chart)],['template','audit',str(chart),'--set','serverUrl=https://api.example.test','--set','workspaceUrl=https://app.example.test','--set','server.env[0].name=FRONT_URL','--set','server.env[0].value=https://app.example.test']]:
    r=subprocess.run([a.helm]+args,capture_output=True,text=True,encoding='utf-8');tests.append({'command':'helm '+' '.join(args).replace(str(chart),'<chart>'),'exit':r.returncode,'stderr':r.stderr.replace(str(td),'<temp>'),'resource_count':r.stdout.count('\nkind:'),'front_url_present':'FRONT_URL' in r.stdout})
   out['helm']=tests
out['effort']={'modernization_pm':[159*.6,159*.8],'modernization_kzt':[159*k*2125*490 for k in [.6,.8]],'rebuild_pm':[383*.6,383*.8],'note':'Arithmetic of reported inputs, not verified actual work.'}
Path(a.out).write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8');print('Wrote',a.out)
